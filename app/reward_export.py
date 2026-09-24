"""Exportación Excel local de premios; no vuelve a calcular snapshots históricos."""
import os
from pathlib import Path
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_rewards(path, data):
    if data.get("legacy"):
        raise ValueError("Este cierre no contiene un snapshot de premios")
    target = Path(path)
    book = Workbook()
    summary = book.active
    summary.title = "RESUMEN"
    detail = book.create_sheet("DETALLE PREMIOS")
    rules = data["rules"]
    state = "CONFIRMADO" if data.get("confirmed") and data.get("week_status")=="CERRADA" else "PRELIMINAR — NO LIQUIDAR"
    for sheet in (summary,detail):
        sheet.append([f"SHES · Comisiones y premios · {data['week_id']} · {state}"])
        sheet.append(["Comisión base: 3% antes de IVA. Premios adicionales acumulativos. Valores vacíos: no disponibles."])
    summary.append(["Vendedor","Venta bruta","Devoluciones","Venta neta antes IVA","Comisión 3%"] + [f"{r['name']} [#{r['id']}]" for r in rules] +
                   ["Otros premios","Total premios","Total comisión + premios","Estado semana","Control final"])
    for seller in data["sellers"]:
        amounts = {r["rule"]["id"]:r["earned"] for r in seller["results"]}
        summary.append([seller["seller"],seller.get("sale_gross"),seller.get("returns",0),seller["sale_net"],seller["base_commission"]] +
                       [amounts.get(r["id"],0) for r in rules] + [0,seller["total_awards"],seller["total_variable"],
                       data["week_status"],"Confirmado" if seller["control_valid"] else "Pendiente"])
    summary.append(["TOTAL",sum((s.get("sale_gross") or 0) for s in data["sellers"]),sum(s.get("returns",0) for s in data["sellers"]),data["sale_net"],data["base_commission"]] +
                   [sum(r["earned"] for s in data["sellers"] for r in s["results"] if r["rule"]["id"]==rule["id"]) for rule in rules] +
                   [0,data["total_awards"],round(data["base_commission"]+data["total_awards"],2) if data["base_commission"] is not None else None,
                    data["week_status"],state])
    detail.append(["Vendedor","ID regla","Premio","Grupo","Nivel","Métrica","Proveedor","Artículo","Operador",
                   "Actual","Objetivo","Progreso %","Falta","Estado","Monto regla","Premio computado","Explicación",
                   "Valor manual","Nota manual","Fecha manual","Observación manual","Invalidaciones","Nota control final",
                   "Vigente desde","Vigente hasta"])
    for seller in data["sellers"]:
        for result in seller["results"]:
            rule, manual = result["rule"], result.get("manual") or {}
            detail.append([seller["seller"],rule["id"],rule["name"],rule["group"],rule["level"],rule["metric"],rule["provider"],
                rule["article"],rule["operator"],result["actual"],rule["target"],result["progress"],result["gap"],result["status"],
                rule["amount"],result["earned"],result["explanation"],float(manual["value"]) if manual else None,
                manual.get("note"),manual.get("date"),manual.get("observation"),
                " | ".join(f"{i['date']}: {i['reason']} {i['note']}" for i in result["invalidations"]),
                (seller.get("control") or {}).get("note"),rule["date_from"],rule["date_to"]])
    for sheet in (summary,detail):
        sheet.freeze_panes = "D4"
        sheet.auto_filter.ref = f"A3:{get_column_letter(sheet.max_column)}{max(3,sheet.max_row-(1 if sheet is summary else 0))}"
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 28
        sheet.row_dimensions[3].height = 32
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type = "s"  # Evita fórmulas introducidas mediante nombres/notas.
                if cell.row == 1:
                    cell.font = Font(name="Calibri",bold=True,size=14,color="D91C28")
                elif cell.row == 3:
                    cell.font = Font(name="Calibri",bold=True,color="302D29")
                    cell.fill = PatternFill("solid",fgColor="FFCF24")
                    cell.alignment = Alignment(wrap_text=True,vertical="center")
                elif cell.row>3:
                    cell.font = Font(name="Calibri",size=11)
                    if cell.row%2==0:
                        cell.fill = PatternFill("solid",fgColor="FFF9E6")
                    if isinstance(cell.value,(int,float)):
                        cell.number_format = '#,##0.00;[Red]-#,##0.00'
        for col in range(1,sheet.max_column+1):
            width = max(len(str(sheet.cell(row,col).value or "")) for row in range(3,sheet.max_row+1))
            sheet.column_dimensions[get_column_letter(col)].width = min(48,max(16,width+2))
        sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=sheet.max_column)
        sheet.merge_cells(start_row=2,start_column=1,end_row=2,end_column=sheet.max_column)
    for cell in summary[summary.max_row]:
        cell.font = Font(name="Calibri",bold=True,color="302D29")
        cell.fill = PatternFill("solid",fgColor="FFCF24")
    # Guardado atómico: una exportación fallida no daña un archivo anterior.
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=target.parent,suffix=".xlsx",delete=False) as stream:
            temporary = Path(stream.name)
        book.save(temporary)
        os.replace(temporary,target)
    finally:
        book.close()
        if temporary and temporary.exists():
            temporary.unlink()
    return target
