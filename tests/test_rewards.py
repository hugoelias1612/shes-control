"""Premios: acumulación, base neta, controles, cierres, exportación y tablas."""
from copy import deepcopy
from datetime import date
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from app.rewards import RewardService, compare, METRICS, metric_key
from app.week_store import WeekStore, file_hash
from app.week_service import WeekService
from app.reward_export import export_rewards


class RewardsTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.store=WeekStore(self.root/'private')
        self.key=self.store.ensure_week(date(2026,9,21))["id"]
        self.service=WeekService(self.store,today=date(2026,9,30));self.engine=RewardService(self.service)
        seller=dict(seller="A",sale=10890000,sale_net=9000000,average_ticket_net=90000,
            coverage_pct=61,conversion_pct=80,buyers=100,logical_orders=130,average_article_mix=4.5,
            reward_availability=dict(orders=True,articles=True,sigo=True),
            providers=[dict(provider="GAONA",sale=968000,sale_net=800000,clients=35,buyer_coverage_pct=35)],
            articles=[dict(code="10",name="Chips",provider="GAONA",sale=121000,sale_net=100000,quantity=40.5,clients=12,client_codes=[str(i) for i in range(12)])])
        self.data=dict(status="EN_CURSO",week_id=self.key,source_revision=0,sellers=[seller])
        self.patch=patch.object(self.service,"metrics",side_effect=lambda _:deepcopy(self.data));self.patch.start();self.addCleanup(self.patch.stop)

    def rule(self,**changes):
        rule=dict(name="Ventas",metric="sale_net",operator=">=",target=7000000,amount=10000,date_from="2026-09-21")
        rule.update(changes);return self.engine.save_rule(rule)

    def result(self):return self.engine.calculate(self.key)

    def control(self,**overrides):
        row=self.result()["sellers"][0]
        values={"sale_net||":row["sale_net"]}
        values.update({metric_key(r["rule"]):r["actual"] for r in row["results"] if not r["rule"]["manual"]})
        values.update(overrides)
        self.engine.confirm_control(self.key,"A",values,"Revisión final documentada","2026-09-30",row["source_signature"])

    def test_levels_accumulate_and_commission_is_separate(self):
        for level,target,amount in [(1,7000000,10000),(2,8000000,5000),(3,9000000,10000)]:
            self.rule(name=f"Nivel {level}",group="Ventas",level=level,target=target,amount=amount)
        data=self.result();seller=data["sellers"][0]
        self.assertEqual(seller["total_awards"],25000)
        self.assertEqual(seller["base_commission"],270000)
        self.assertEqual(seller["total_variable"],295000)
        self.assertEqual([r["status"] for r in seller["results"]],["CUMPLIDO_PRELIMINAR"]*3)
        self.assertGreater(seller["results"][0]["progress"],100)

    def test_all_metrics_and_provider_penetration(self):
        for metric,(_,_,source,_) in METRICS.items():
            if source=="manual":continue
            self.rule(name=metric,metric=metric,target=1,provider="GAONA" if source in {"provider","article"} else "",article="10" if source=="article" else "")
        rows=self.result()["sellers"][0]["results"]
        self.assertEqual(len(rows),13)
        self.assertTrue(all(r["status"]=="CUMPLIDO_PRELIMINAR" for r in rows))
        values={r["rule"]["metric"]:r["actual"] for r in rows}
        self.assertEqual(values["provider_coverage"],35)
        self.assertEqual(values["provider_sale"],800000)
        self.assertEqual(values["article_quantity"],40.5)
        self.assertEqual(values["article_buyers"],12)

    def test_operators_and_explanations(self):
        for op,value,reached in [(">=",10,True),(">",10,False),("<=",9,True),("<",10,False),("=",10,True),("=",11,False)]:
            rule=dict(metric="buyers",target=10,operator=op)
            self.assertEqual(compare(value,rule)[0],reached)
            self.assertTrue(compare(value,rule)[3])
        self.assertIn("puntos porcentuales",compare(47.6,dict(metric="coverage",target=55,operator=">="))[3])
        self.assertAlmostEqual(compare(7900000,dict(metric="sale_net",target=8500000,operator=">="))[2],600000)
        for op in [">=",">","<=","<","="]:
            self.assertIsInstance(compare(0,dict(metric="buyers",target=0,operator=op))[1],float)

    def test_missing_data_is_not_zero_or_an_earned_award(self):
        self.rule(metric="coverage",operator="<=",target=10)
        self.data["sellers"][0]["reward_availability"]["sigo"]=False
        row=self.result()["sellers"][0]["results"][0]
        self.assertIsNone(row["actual"]);self.assertEqual(row["earned"],0);self.assertEqual(row["status"],"PENDIENTE")
        self.data["sellers"][0]["sale_net"]=None
        self.assertIsNone(self.result()["base_commission"])
        self.assertIsNone(self.result()["cost_pct"])

    def test_manual_values_and_notes_are_audited(self):
        rid=self.rule(metric="reactivated_clients",target=5)
        row=self.result()["sellers"][0]["results"][0]
        self.assertEqual(row["status"],"PENDIENTE_CONTROL_MANUAL")
        self.engine.manual(self.key,"A",rid,6,"Kiosco Pepe\nEl Sol","2026-09-29","Verificados")
        row=self.result()["sellers"][0]["results"][0]
        self.assertEqual(row["status"],"CUMPLIDO_PRELIMINAR")
        self.assertIn("Kiosco Pepe",row["manual"]["note"])
        with self.assertRaises(ValueError):self.engine.manual(self.key,"A",rid,1.5,"","2026-09-29")
        self.assertTrue(self.store.query("SELECT * FROM reward_audit WHERE action='valor_manual'"))

    def test_individual_and_all_invalidations_and_reversal(self):
        first=self.rule();second=self.rule(name="Otro")
        with self.assertRaises(ValueError):self.engine.invalidate(self.key,"A",first,"","2026-09-29")
        self.engine.invalidate(self.key,"A",first,"Devolución","2026-09-29","Nota")
        self.assertEqual(self.result()["total_awards"],10000)
        self.engine.invalidate(self.key,"A",None,"Control administrativo","2026-09-29")
        self.assertEqual(self.result()["total_awards"],0)
        for row in self.store.query("SELECT id FROM reward_invalidations"):
            self.engine.revoke_invalidation(self.key,row["id"],"Control resuelto")
        self.assertEqual(self.result()["total_awards"],20000)

    def test_final_recalculation_drops_preliminary_level(self):
        self.rule();self.rule(name="Nivel 2",target=8500000,amount=5000)
        self.assertEqual(self.result()["total_awards"],15000)
        with self.assertRaisesRegex(ValueError,"control final"):self.engine.closing(self.key,self.data)
        self.control(**{"sale_net||":7900000})
        final=self.engine.closing(self.key,self.data)
        self.assertEqual(final["sellers"][0]["base_commission"],237000)
        self.assertEqual(final["total_awards"],10000)
        self.assertEqual([r["status"] for r in final["sellers"][0]["results"]],["CONFIRMADO","NO_ALCANZADO"])

    def test_source_change_or_rule_change_invalidates_final_control(self):
        rid=self.rule();self.control()
        self.assertTrue(self.result()["sellers"][0]["control_valid"])
        rule=self.engine.rules()[0];rule["target"]=9999999;self.engine.save_rule(rule,rid,rule["version"])
        self.assertFalse(self.result()["sellers"][0]["control_valid"])
        self.control();self.data["source_revision"]+=1
        self.assertFalse(self.result()["sellers"][0]["control_valid"])

    def test_manual_change_invalidates_control(self):
        rid=self.rule(metric="new_clients",target=2)
        self.engine.manual(self.key,"A",rid,3,"Chequeo","2026-09-29");self.control()
        self.engine.manual(self.key,"A",rid,1,"Corrección","2026-09-30")
        self.assertFalse(self.result()["sellers"][0]["control_valid"])

    def test_scope_dates_inactive_archive_and_rule_validation(self):
        self.rule(scope=["B"]);self.rule(active=False);self.rule(date_from="2026-09-28")
        self.rule(date_from="2026-09-01",date_to="2026-09-20")
        rid=self.rule();self.engine.archive(rid)
        self.assertEqual(self.result()["sellers"][0]["results"],[])
        for changes in [dict(amount=-1),dict(target=float('nan')),dict(operator="AND"),dict(metric="provider_sale"),dict(metric="article_quantity"),dict(metric="buyers",target=1.5)]:
            with self.assertRaises(ValueError):self.rule(**changes)

    def test_seven_percent_alert_does_not_cap_awards(self):
        self.rule(amount=900000)
        data=self.result()
        self.assertTrue(data["over_seven_percent"])
        self.assertEqual(data["total_awards"],900000)
        self.assertEqual(data["cost_pct"],13)
        self.data["sellers"][0]["sale_net"]=0
        self.assertIsNone(self.result()["cost_pct"])

    def test_snapshot_is_immutable_after_edit_and_reopen(self):
        rid=self.rule();self.control()
        awards=self.engine.closing(self.key,self.data)
        data=dict(self.data,awards=awards,status="CERRADA")
        path=self.store.save_snapshot(self.key,data,kind="cierre",close=True);digest=file_hash(path)
        self.data=data
        rule=self.engine.rules()[0];rule["amount"]=1;self.engine.save_rule(rule,rid,rule["version"])
        self.assertEqual(self.result()["total_awards"],10000)
        with self.assertRaises(ValueError):self.engine.invalidate(self.key,"A",None,"X","2026-09-30")
        self.store.reopen(self.key,"Revisión de ajustes")
        self.data["status"]="REABIERTA";self.data["source_revision"]=self.store.week(self.key)["revision"]
        self.assertFalse(self.result()["sellers"][0]["control_valid"])
        self.assertEqual(file_hash(path),digest)
        self.assertEqual(json.loads(path.read_text())["awards"]["total_awards"],10000)

    def test_persist_results_and_restart(self):
        self.rule();self.engine.calculate(self.key,persist=True)
        restarted=WeekStore(self.store.root)
        self.assertEqual(len(restarted.query("SELECT * FROM reward_results")),1)
        self.assertEqual(len(restarted.query("SELECT * FROM reward_rules")),1)

    def test_excel_all_sellers_columns_numeric_and_formula_injection(self):
        from openpyxl import load_workbook
        self.rule(name="=1+1")
        second=deepcopy(self.data["sellers"][0]);second["seller"]="=HYPERLINK(1)";self.data["sellers"].append(second)
        path=self.root/'premios.xlsx';export_rewards(path,self.result())
        book=load_workbook(path)
        self.assertEqual(book.sheetnames,["RESUMEN","DETALLE PREMIOS"])
        summary=book["RESUMEN"]
        self.assertEqual(summary.max_row,6)
        self.assertEqual(summary['C4'].value,270000)
        self.assertEqual(summary['A5'].data_type,'s')
        self.assertIn("PRELIMINAR",summary['A1'].value)
        self.assertEqual(summary.freeze_panes,"D4")
        self.assertEqual(summary["A6"].value,"TOTAL")
        self.assertEqual(summary["C6"].value,540000)
        book.close()

    def test_tables_allow_resize_move_and_restore_selection(self):
        os.environ['QT_QPA_PLATFORM']='offscreen'
        from PySide6.QtWidgets import QApplication,QWidget,QVBoxLayout,QHeaderView
        from app.week_window import table
        from app.tables import capture_tables,restore_tables
        app=QApplication.instance() or QApplication([])
        parent=QWidget();layout=QVBoxLayout(parent);grid=table(["Vendedor","Venta"],[["A",100],["B",2]]);layout.addWidget(grid)
        self.assertEqual(grid.horizontalHeader().sectionResizeMode(0),QHeaderView.Interactive)
        self.assertTrue(grid.horizontalHeader().sectionsMovable());self.assertFalse(grid.horizontalHeader().stretchLastSection())
        grid.setColumnWidth(0,350);grid.selectRow(0);saved=capture_tables(parent);grid.clearSelection();grid.setColumnWidth(0,100)
        restore_tables(parent,saved);self.assertEqual(grid.columnWidth(0),350);self.assertEqual(len(grid.selectionModel().selectedRows()),1)
        parent.close();app.processEvents()

    def test_configuration_ui_and_detail_can_be_opened(self):
        os.environ['QT_QPA_PLATFORM']='offscreen'
        from PySide6.QtWidgets import QApplication
        from app.reward_window import RewardsPanel,RuleDialog,RewardDetail
        app=QApplication.instance() or QApplication([])
        self.rule();panel=RewardsPanel(self.service,self.key,self.data)
        self.assertEqual(panel.tabs.count(),3)
        rule=RuleDialog(self.engine,self.key,self.data);self.assertEqual(rule.metric_combo.currentData(),"sale_net")
        detail=RewardDetail(panel,"A");self.assertEqual(detail.grid.rowCount(),1)
        rule.show();panel.show();detail.show();app.processEvents()
        for w in [detail,rule,panel]:w.close()
        app.processEvents()

    def test_rule_dialog_saves_duplicates_and_manual_metric_is_forced(self):
        os.environ['QT_QPA_PLATFORM']='offscreen'
        from PySide6.QtWidgets import QApplication,QDialogButtonBox,QDialog
        from app.reward_window import RuleDialog
        app=QApplication.instance() or QApplication([])
        dialog=RuleDialog(self.engine,self.key,self.data)
        dialog.name.setText("Nuevos clientes")
        dialog.metric_combo.setCurrentIndex(dialog.metric_combo.findData("new_clients"))
        self.assertTrue(dialog.manual.isChecked());self.assertFalse(dialog.manual.isEnabled())
        dialog.target.setValue(5);dialog.amount.setValue(10000)
        dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Save).click()
        self.assertEqual(dialog.result(),QDialog.Accepted)
        rule=self.engine.rules()[0]
        duplicate=RuleDialog(self.engine,self.key,self.data,rule,True)
        duplicate.findChild(QDialogButtonBox).button(QDialogButtonBox.Save).click()
        self.assertEqual(len(self.engine.rules()),2)
        self.assertNotEqual(self.engine.rules()[0]["id"],self.engine.rules()[1]["id"])
        dialog.close();duplicate.close();app.processEvents()

    def test_stale_rule_editor_and_snapshot_inputs_are_rejected(self):
        rid=self.rule();old=self.engine.rules()[0]
        self.engine.save_rule(dict(old,amount=20000),rid,old["version"])
        with self.assertRaises(ValueError):self.engine.save_rule(old,rid,old["version"])
        data=dict(self.data,awards=self.result())
        self.rule(name="Nueva regla")
        with self.assertRaisesRegex(ValueError,"premios cambiaron"):
            self.store.save_snapshot(self.key,data,kind="cierre",close=True)
        self.assertNotEqual(self.store.week(self.key)["status"],"CERRADA")


class RewardIntegrationTests(unittest.TestCase):
    def test_net_columns_reassignment_articles_and_existing_close(self):
        import test_weeks
        fixture=test_weeks.WeekTests();fixture.setUp();self.addCleanup(fixture.doCleanups)
        order=fixture.order(total=121)
        order.update({"NETO GRAVADO":80,"NO GRAVADO":20})
        line=fixture.line(amount=121);line["Importes Netos"]=100
        fixture.upload(fixture.points(),fixture.points(branch="resistencia",seller="B",client=2),fixture.report([order]),fixture.report([line]))
        data=fixture.service.metrics(fixture.key)
        self.assertEqual(data["company"]["sale"],121)
        self.assertEqual(data["company"]["sale_net"],100)
        self.assertEqual(data["articles"][0]["sale_net"],100)
        self.assertEqual(data["providers"][0]["sale_net"],100)
        session,_,_,rev=fixture.service.build_session(fixture.key);session.orders[0].assigned_seller="B"
        fixture.service.confirm_review(fixture.key,session,rev)
        data=fixture.service.metrics(fixture.key)
        self.assertEqual(next(s for s in data["sellers"] if s["seller"]=="B")["sale_net"],100)
        engine=RewardService(fixture.service)
        engine.save_rule(dict(name="Venta",metric="sale_net",target=100,amount=10,date_from="2026-09-21"))
        from datetime import timedelta
        for offset in range(1,6):fixture.store.set_worked(fixture.key,test_weeks.MON+timedelta(days=offset),test_weeks.BRANCHES,False,"Feriado")
        session,_,_,rev=fixture.service.build_session(fixture.key);fixture.service.confirm_review(fixture.key,session,rev)
        for day in [test_weeks.MON,test_weeks.SUN]:fixture.store.confirm_liquidations(fixture.key,day,"Sin pendientes")
        with self.assertRaisesRegex(ValueError,"control final"):fixture.service.close(fixture.key)
        for seller in engine.calculate(fixture.key)["sellers"]:
            engine.confirm_control(fixture.key,seller["seller"],{"sale_net||":seller["sale_net"]},"Verificado","2026-09-30",seller["source_signature"])
        path=fixture.service.close(fixture.key)
        saved=json.loads(path.read_text())
        self.assertEqual(saved["awards"]["total_awards"],10)
        self.assertTrue(saved["awards"]["confirmed"])
        self.assertEqual(engine.calculate(fixture.key)["total_awards"],10)
